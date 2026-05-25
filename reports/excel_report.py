import os
import sys
import json
import sqlite3
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Add project root to sys.path
project_root = Path(__file__).resolve().parent.parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from core.logger import logger

def generate_excel_report(batch_id: str, db_path="data/durian.db", output_dir="reports/outputs"):
    logger.info(f"Generating Excel report for batch: {batch_id}")
    
    # Create output directory
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    file_path = out_dir / f"batch_report_{batch_id}.xlsx"

    # Query DB
    if not os.path.exists(db_path):
        logger.error(f"Database {db_path} not found. Cannot generate report.")
        return None

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    try:
        # Fetch batch info
        cursor.execute("SELECT * FROM batches WHERE batch_id = ?;", (batch_id,))
        batch_row = cursor.fetchone()
        if not batch_row:
            logger.error(f"Batch {batch_id} not found in DB.")
            return None
        batch = dict(batch_row)

        # Fetch detections
        cursor.execute("SELECT * FROM detections WHERE batch_id = ? ORDER BY timestamp ASC;", (batch_id,))
        detections_rows = cursor.fetchall()
        detections = [dict(r) for r in detections_rows]

        # Initialize Workbook
        wb = openpyxl.Workbook()
        
        # Styles
        font_family = "Segoe UI"
        title_font = Font(name=font_family, size=16, bold=True, color="2C3E50")
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        bold_font = Font(name=font_family, size=10, bold=True, color="2C3E50")
        normal_font = Font(name=font_family, size=10, color="333333")
        
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        summary_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
        
        thin_side = Side(border_style="thin", color="BDC3C7")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # ----------------------------------------------------
        # SHEET 1: Summary Sheet
        # ----------------------------------------------------
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.views.sheetView[0].showGridLines = True
        
        ws1["A1"] = f"BÁO CÁO PHÂN LOẠI SẦU RIÊNG - BATCH {batch_id}"
        ws1["A1"].font = title_font
        ws1.merge_cells("A1:D1")
        ws1.row_dimensions[1].height = 40
        
        # Metadata
        metadata = [
            ("Mã Lô (Batch ID):", batch_id),
            ("Thời Gian Bắt Đầu:", batch.get("start_time", "N/A")),
            ("Thời Gian Kết Thúc:", batch.get("end_time", "N/A")),
            ("Tổng Số Lượng Quả:", batch.get("total_count", 0)),
        ]
        
        for idx, (label, val) in enumerate(metadata, 3):
            ws1[f"A{idx}"] = label
            ws1[f"A{idx}"].font = bold_font
            ws1[f"B{idx}"] = val
            ws1[f"B{idx}"].font = normal_font
            ws1.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=4)
            ws1.row_dimensions[idx].height = 20

        # Grade Distribution table header
        start_row = 9
        ws1.cell(row=start_row, column=1, value="Phân Loại (Grade)").font = header_font
        ws1.cell(row=start_row, column=1).fill = header_fill
        ws1.cell(row=start_row, column=1).alignment = align_center
        
        ws1.cell(row=start_row, column=2, value="Số Lượng (Count)").font = header_font
        ws1.cell(row=start_row, column=2).fill = header_fill
        ws1.cell(row=start_row, column=2).alignment = align_center
        
        ws1.cell(row=start_row, column=3, value="Tỷ Lệ (%)").font = header_font
        ws1.cell(row=start_row, column=3).fill = header_fill
        ws1.cell(row=start_row, column=3).alignment = align_center
        
        ws1.row_dimensions[start_row].height = 25
        
        grades_data = [
            ("Hạng A (Grade A)", batch.get("grade_a", 0)),
            ("Hạng B (Grade B)", batch.get("grade_b", 0)),
            ("Hạng C (Grade C)", batch.get("grade_c", 0)),
            ("Loại (Reject)", batch.get("reject", 0))
        ]
        
        total = max(1, batch.get("total_count", 0))
        for g_idx, (g_name, g_count) in enumerate(grades_data, start_row + 1):
            ws1.cell(row=g_idx, column=1, value=g_name).font = bold_font
            ws1.cell(row=g_idx, column=1).border = border_all
            ws1.cell(row=g_idx, column=1).alignment = align_left
            
            ws1.cell(row=g_idx, column=2, value=g_count).font = normal_font
            ws1.cell(row=g_idx, column=2).border = border_all
            ws1.cell(row=g_idx, column=2).alignment = align_right
            
            ratio = g_count / total
            cell_ratio = ws1.cell(row=g_idx, column=3, value=ratio)
            cell_ratio.font = normal_font
            cell_ratio.border = border_all
            cell_ratio.number_format = '0.0%'
            cell_ratio.alignment = align_right
            
            ws1.row_dimensions[g_idx].height = 20

        # ----------------------------------------------------
        # SHEET 2: Detections Log Sheet
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="Detections Log")
        ws2.views.sheetView[0].showGridLines = True
        
        headers = ["ID", "Thời Gian", "Kết Quả Phân Loại", "Tổng Số Lỗi", "Mức Độ Tin Cậy", "Đường Dẫn Ảnh"]
        ws2.row_dimensions[1].height = 25
        
        for col_idx, text in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col_idx, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            
        for row_idx, det in enumerate(detections, 2):
            ws2.row_dimensions[row_idx].height = 20
            
            ws2.cell(row=row_idx, column=1, value=det["id"]).font = normal_font
            ws2.cell(row=row_idx, column=1).border = border_all
            ws2.cell(row=row_idx, column=1).alignment = align_center
            
            ws2.cell(row=row_idx, column=2, value=det["timestamp"]).font = normal_font
            ws2.cell(row=row_idx, column=2).border = border_all
            ws2.cell(row=row_idx, column=2).alignment = align_center
            
            ws2.cell(row=row_idx, column=3, value=det["final_grade"]).font = bold_font
            ws2.cell(row=row_idx, column=3).border = border_all
            ws2.cell(row=row_idx, column=3).alignment = align_center
            
            ws2.cell(row=row_idx, column=4, value=det["defect_count"]).font = normal_font
            ws2.cell(row=row_idx, column=4).border = border_all
            ws2.cell(row=row_idx, column=4).alignment = align_right
            
            cell_conf = ws2.cell(row=row_idx, column=5, value=det["confidence"])
            cell_conf.font = normal_font
            cell_conf.border = border_all
            cell_conf.number_format = '0.0%'
            cell_conf.alignment = align_right
            
            ws2.cell(row=row_idx, column=6, value=det["image_path"] or "").font = normal_font
            ws2.cell(row=row_idx, column=6).border = border_all
            ws2.cell(row=row_idx, column=6).alignment = align_left

        # ----------------------------------------------------
        # SHEET 3: Defect Frequencies
        # ----------------------------------------------------
        ws3 = wb.create_sheet(title="Defect Frequencies")
        ws3.views.sheetView[0].showGridLines = True
        
        ws3.row_dimensions[1].height = 25
        ws3.cell(row=1, column=1, value="Loại Khuyết Tật (Defect Type)").font = header_font
        ws3.cell(row=1, column=1).fill = header_fill
        ws3.cell(row=1, column=1).alignment = align_center
        
        ws3.cell(row=1, column=2, value="Tần Suất Xuất Hiện (Occurrences)").font = header_font
        ws3.cell(row=1, column=2).fill = header_fill
        ws3.cell(row=1, column=2).alignment = align_center

        # Aggregate defect counts from detections
        defect_counts = {'crack': 0, 'dark_spot': 0, 'fungus': 0, 'thorn_split': 0, 'reject': 0}
        for det in detections:
            try:
                labels = json.loads(det["defect_types"]) if det["defect_types"] else []
                for label in labels:
                    if label in defect_counts:
                        defect_counts[label] += 1
            except Exception:
                pass
                
        for row_idx, (label, count) in enumerate(defect_counts.items(), 2):
            ws3.row_dimensions[row_idx].height = 20
            
            ws3.cell(row=row_idx, column=1, value=label).font = bold_font
            ws3.cell(row=row_idx, column=1).border = border_all
            ws3.cell(row=row_idx, column=1).alignment = align_left
            
            ws3.cell(row=row_idx, column=2, value=count).font = normal_font
            ws3.cell(row=row_idx, column=2).border = border_all
            ws3.cell(row=row_idx, column=2).alignment = align_right

        # Auto-fit columns across all sheets
        for sheet in [ws1, ws2, ws3]:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    # Ignore merged cells in length calculation
                    if cell.coordinate in sheet.merged_cells:
                        continue
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(file_path)
        logger.info(f"Excel report saved to: {file_path}")
        return str(file_path)
    except Exception as e:
        logger.error(f"Error compiling Excel report: {e}")
        return None
    finally:
        conn.close()

if __name__ == "__main__":
    if len(sys.argv) > 1:
        generate_excel_report(sys.argv[1])
    else:
        print("Usage: python excel_report.py <batch_id>")
